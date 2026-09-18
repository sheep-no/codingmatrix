"""PPT 模板/修改意图回归测试。

覆盖 docs/evolution/modules/pptx_toolkit.md 中核实的缺陷：
- PPX1：自定义模板上传端点调用不存在的 parser.parse，且对 dataclass 直接 json.dump
- PPX2：主题色/字体应用因 master.theme 不存在而恒失败（且 clrScheme 需在 themeElements 下查找）
- PPX4：字号修改意图 property_value 恒 None，执行端永不触发
"""
import io
import json
import os
import sys
import tempfile

import pytest
from fastapi.testclient import TestClient
from lxml import etree
from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Pt

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))))

from app.main import app
from app.utils.security import verify_token
from app.utils.pptx.custom_template import TemplateConverter
from app.utils.pptx.modify_intent_parser import ModifyIntentParser
from app.utils.pptx.ppt_modifier import PPTModifier
from app.utils.pptx.templates.base import TemplateCategory, TemplateConfig


def _save_minimal_pptx_bytes() -> bytes:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    if slide.shapes.title:
        slide.shapes.title.text = "标题"
    for shape in slide.shapes:
        if shape.has_text_frame and shape != slide.shapes.title:
            shape.text_frame.paragraphs[0].text = "正文内容"
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def _theme_element(prs):
    theme_part = next(
        rel.target_part
        for rel in prs.slide_masters[0].part.rels.values()
        if rel.reltype.endswith("/theme")
    )
    return etree.fromstring(theme_part.blob)


def _sample_config() -> TemplateConfig:
    return TemplateConfig(
        template_id="t",
        name="n",
        name_zh="n",
        category=TemplateCategory.BUSINESS,
        description="d",
        primary_color="1F4E79",
        text_color="333333",
        secondary_color="2E75B6",
        accent_color="70AD47",
        light_text_color="666666",
        background_color="FFFFFF",
        title_font_en="Arial",
        title_font="微软雅黑",
        body_font_en="Calibri",
        body_font="宋体",
    )


class TestCustomTemplateUpload:
    """PPX1：上传端点应返回可用的解析配置"""

    @pytest.fixture(autouse=True)
    def _auth_and_cwd(self, tmp_path, monkeypatch):
        app.dependency_overrides[verify_token] = lambda: {"sub": "u1", "role": "user"}
        monkeypatch.chdir(tmp_path)
        yield tmp_path
        app.dependency_overrides.clear()

    def _upload(self, client):
        return client.post(
            "/api/v1/pptx/templates/upload",
            files={"file": ("t.pptx", _save_minimal_pptx_bytes(), "application/octet-stream")},
            data={"name": "我的模板", "description": "desc"},
        )

    def test_upload_returns_parsed_config(self):
        resp = self._upload(TestClient(app))

        assert resp.status_code == 200
        body = resp.json()
        assert body["config"] is not None, body.get("message")
        assert body["config"]["category"] == "business"
        assert "primary_color" in body["config"]

    def test_upload_writes_json_config_file(self, tmp_path):
        resp = self._upload(TestClient(app))

        template_id = resp.json()["template_id"]
        config_path = tmp_path / "configs" / "ppt" / "custom_templates" / f"{template_id}.json"
        assert config_path.exists()
        saved = json.loads(config_path.read_text(encoding="utf-8"))
        assert saved["template_id"] == "custom_user_upload"

    def test_to_dict_is_json_serializable_with_enum_values(self):
        config = _sample_config()
        config.layouts = {TemplateCategory.BUSINESS: {"recommended_type": "title_content"}}
        # asdict 保留枚举键，to_dict 需转为字符串，否则 json.dumps 抛 TypeError
        dumped = json.dumps(config.to_dict())
        assert "business" in dumped


class TestThemeApplication:
    """PPX2：主题色与字体应真正写入主题 part"""

    def _apply_and_reload(self):
        prs = Presentation()
        TemplateConverter().apply_config_to_presentation(_sample_config(), prs)
        buf = io.BytesIO()
        prs.save(buf)
        buf.seek(0)
        return Presentation(buf)

    def test_theme_colors_written_to_clr_scheme(self):
        theme_el = _theme_element(self._apply_and_reload())
        clr_scheme = theme_el.find(".//" + qn("a:clrScheme"))

        def _color(slot):
            node = clr_scheme.find(qn("a:" + slot))
            assert node is not None and len(node), f"{slot} 未写入颜色"
            return node[0].get("val")

        assert _color("dk1") == "1F4E79"
        assert _color("accent3") == "FFFFFF"

    def test_theme_fonts_written_to_font_scheme(self):
        theme_el = _theme_element(self._apply_and_reload())
        font_scheme = theme_el.find(".//" + qn("a:fontScheme"))

        major = font_scheme.find(qn("a:majorFont"))
        minor = font_scheme.find(qn("a:minorFont"))
        assert major.find(qn("a:latin")).get("typeface") == "Arial"
        assert major.find(qn("a:ea")).get("typeface") == "微软雅黑"
        assert minor.find(qn("a:latin")).get("typeface") == "Calibri"
        assert minor.find(qn("a:ea")).get("typeface") == "宋体"


class TestSizeIntent:
    """PPX4：字号修改意图应携带具体值并被执行"""

    def test_size_intent_carries_value(self):
        intent = ModifyIntentParser().parse("把正文字号改成24")
        size_targets = [t for t in intent.targets if t.property_name == "size"]
        assert size_targets
        assert size_targets[0].property_value == "24"

    def test_size_intent_applied_end_to_end(self, tmp_path):
        src = tmp_path / "in.pptx"
        dst = tmp_path / "out.pptx"
        src.write_bytes(_save_minimal_pptx_bytes())

        intent = ModifyIntentParser().parse("把第1页字号改成24")
        assert PPTModifier(str(src)).apply_modifications(intent, str(dst)) is True

        prs = Presentation(str(dst))
        sizes = [
            para.font.size
            for shape in prs.slides[0].shapes
            if shape.has_text_frame
            for para in shape.text_frame.paragraphs
        ]
        assert Pt(24) in sizes
