from types import SimpleNamespace
import shutil
from unittest.mock import AsyncMock, Mock

import pytest
from pptx import Presentation

from app.api.v1 import aiGeneratorPptx as api
from app.schema.ppt_outline import SlideRegenerateRequest
from pydantic import ValidationError
from app.services.ppt_quality_orchestrator import review_rendered_deck, run_quality_pipeline
from app.services.ppt_generation_persistence import serialize_quality_report
from app.utils import vision


def test_single_slide_request_requires_edit_and_base_version():
    with pytest.raises(ValidationError):
        SlideRegenerateRequest()


@pytest.mark.asyncio
async def test_reflow_preserves_approved_content_over_capacity():
    blocks = [{"content": f"原文 {index}"} for index in range(8)]
    pages, report = await run_quality_pipeline([{
        "id": "retained", "preserve_content": True, "content_blocks": blocks, "capacity": {"max_items": 3},
    }])
    assert pages[0]["content_blocks"] == blocks
    assert "retained" in report.manual_review_slides


@pytest.mark.asyncio
async def test_approved_duplicate_pages_survive_normalization_and_export(tmp_path):
    from copy import deepcopy

    outline = {"title": "报告", "status": "approved", "slides": [
        {"id": f"page-{index}", "position": index, "title": "重复页面",
         "key_message": "保留结论", "slide_type": "key_points", "narrative_role": "opportunity_map",
         "content_blocks": [{"type": "text", "content": "保留原文", "metadata": {}}]}
        for index in range(3)
    ]}
    revised = deepcopy(outline)
    revised["slides"][1]["title"] = "修改目标页"
    revised["slides"][1]["content_blocks"][0]["content"] = "用户修改后的正文"
    outputs = []
    for index, source in enumerate([outline, revised]):
        normalized = api._normalize_approved_outline(api._normalize_approved_outline(source))
        assert [page["id"] for page in normalized["slides"]] == ["page-0", "page-1", "page-2"]
        assert normalized["slides"][2]["content_blocks"] == outline["slides"][2]["content_blocks"]
        path = tmp_path / f"revision-{index}.pptx"
        await api.generate_pptx_file_enhanced(path, normalized, api.PPTGenerationRequest(
            topic="报告", slide_count=4, options={"auto_images": False, "enable_animation": False},
        ))
        deck = Presentation(path)
        outputs.append(["\n".join(shape.text for shape in page.shapes if shape.has_text_frame) for page in deck.slides])
    assert len(outputs[0]) == len(outputs[1]) == 4
    assert outputs[0][0:2] == outputs[1][0:2]
    assert outputs[0][3] == outputs[1][3]
    assert "用户修改后的正文" in outputs[1][2]
    assert outputs[0][2] != outputs[1][2]


@pytest.mark.asyncio
@pytest.mark.parametrize("animation", [True, False])
@pytest.mark.parametrize("auto_images", [True, False])
async def test_generation_respects_images_and_animation(tmp_path, monkeypatch, animation, auto_images):
    images = AsyncMock(return_value=None)
    visual = AsyncMock(return_value=None)
    monkeypatch.setattr(api.image_manager, "get_image_for_slide", images)
    monkeypatch.setattr(api.visual_analyzer, "analyze_ppt_content", visual)
    request = api.PPTGenerationRequest(
        topic="报告", slide_count=2, api_key_token="user-token",
        options={"auto_images": auto_images, "enable_animation": animation},
    )
    path = tmp_path / "deck.pptx"
    await api.generate_pptx_file_enhanced(path, {"title": "报告", "slides": [{
        "id": "slide-1", "title": "机会", "content": ["增长"],
        "narrative_role": "opportunity_map", "asset_intent": {"description": "配图"},
    }]}, request)
    assert images.await_count == int(auto_images)
    assert visual.await_count == int(auto_images)
    deck = Presentation(path)
    assert len(deck.slides) == 2
    for slide in deck.slides:
        transitions = slide._element.xpath("./p:transition/p:fade")
        assert bool(transitions) is animation


@pytest.mark.asyncio
async def test_missing_reviewer_is_reported_as_degraded():
    _, report = await run_quality_pipeline([], "refined")
    assert serialize_quality_report(report)["degraded_stage"] == "vision_review_unavailable"


@pytest.mark.asyncio
async def test_missing_user_credentials_skips_rendering_and_reports_degradation(tmp_path, monkeypatch):
    convert = AsyncMock()
    monkeypatch.setattr(api, "_convert_pptx_to_pdf", convert)
    _, report = await run_quality_pipeline([])
    await review_rendered_deck(tmp_path / "deck.pptx", [], report, None, "1")
    convert.assert_not_awaited()
    assert serialize_quality_report(report)["degraded_stage"] == "vision_review_unavailable"


@pytest.mark.asyncio
@pytest.mark.parametrize("response", [
    '{"issues":[{"issue_type":"text_overflow","severity":"high","message":"裁切","confidence":0.6}]}',
    'invalid json',
])
async def test_visual_review_uses_rendered_pages_and_user_token(tmp_path, monkeypatch, response):
    import subprocess
    from app.utils import vision

    path = tmp_path / "deck.pptx"
    path.touch()
    convert = AsyncMock()
    render = Mock(return_value=SimpleNamespace(stdout=b"\x89PNGimage"))
    call = AsyncMock(return_value=response)
    monkeypatch.setattr(api, "_convert_pptx_to_pdf", convert)
    monkeypatch.setattr(subprocess, "run", render)
    monkeypatch.setattr(vision, "_call_vision_model", call)
    slides, report = await run_quality_pipeline([{"id": "slide-1", "elements": []}])
    await review_rendered_deck(path, slides, report, "user-token", "1")
    convert.assert_awaited_once()
    assert call.call_args.kwargs["api_key_token"] == "user-token"
    assert call.call_args.kwargs["image_base64"].startswith("data:image/png;base64,")
    quality = serialize_quality_report(report)
    if response == 'invalid json':
        assert quality["degraded_stage"] == "vision_review_unavailable"
    else:
        assert quality["degraded_stage"] is None
        assert [args.args[0][2] for args in render.call_args_list] == ["1", "2"]
        assert any(issue["slide_id"] == "slide-1" for issue in quality["issues"])
        assert any(issue["issue_type"] == "vision_review_low_confidence" for issue in quality["issues"])


@pytest.mark.asyncio
async def test_vision_call_accepts_fake_llm_caller(tmp_path):
    image = tmp_path / "sample.png"
    image.write_bytes(b"png")
    fake = AsyncMock(return_value={"choices": [{"message": {"content": "可验证"}}]})
    result = await vision._call_vision_model(
        "data:image/png;base64,cG5n", "描述", "fake-model", vision.Timeout(1),
        llm_caller=fake,
    )
    assert result == "可验证"
    assert fake.await_count == 1
    assert fake.call_args.kwargs["messages"][0]["content"][0]["type"] == "image_url"


@pytest.mark.asyncio
@pytest.mark.skipif(not shutil.which("libreoffice") or not shutil.which("pdftoppm"), reason="PDF rendering tools unavailable")
async def test_real_pptx_pdf_png_rendering_with_mocked_vision(tmp_path, monkeypatch):
    from app.utils import vision

    path = tmp_path / "actual-deck.pptx"
    deck = Presentation()
    deck.slides.add_slide(deck.slide_layouts[6])
    deck.slides.add_slide(deck.slide_layouts[6])
    deck.save(path)
    call = AsyncMock(return_value='{"issues": []}')
    monkeypatch.setattr(vision, "_call_vision_model", call)
    slides, report = await run_quality_pipeline([{"id": "slide-1", "elements": []}])
    await review_rendered_deck(path, slides, report, "mock-user-token", "1")
    assert path.with_suffix(".pdf").is_file()
    assert call.await_count == 2
    assert len(call.call_args.kwargs["image_base64"]) > 1000
    assert serialize_quality_report(report)["degraded_stage"] is None
